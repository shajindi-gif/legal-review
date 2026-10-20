"""报告导出（Word .docx）测试 —— R1-1。

两层：
1. 纯函数层：markdown_to_docx 产出合法 docx，标题/加粗/列表/表格可还原。
2. HTTP 层：GET /tasks/{id}/report/export 返回 docx 字节流 + 正确 header；
   报告未生成 → 404；鉴权/多租户复用 _load_task（admin 可见任意任务）。
"""

from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

from fastapi import FastAPI, Request
from fastapi.responses import JSONResponse
from httpx import ASGITransport, AsyncClient

from app.api.deps import TenantContext, get_db, get_tenant_context
from app.api.v1.tasks import router as tasks_router
from app.core.errors import AppError, error_response
from app.services.core_auth_client import CorePrincipal
from app.services.report_export import markdown_to_docx

SAMPLE_MD = """# 审查意见书

## 一、文件基本情况
本文件为 **XX县关于促进中小企业发展的若干意见**。

## 二、审查依据
- 《中小企业促进法》第十条
- 《优化营商环境条例》

## 三、发现问题
1. 主体合法性：制定机关适格
2. 程序完整性：缺少公开征求意见

## 四、风险等级
| 项目 | 结论 |
| --- | --- |
| 主体 | PASS |
| 程序 | RISK |
"""


# ---------- 纯函数层 ----------


def test_markdown_to_docx_returns_valid_docx() -> None:
    """产出为合法 docx（PK 压缩包签名），可被 python-docx 重新打开。"""
    from docx import Document

    data = markdown_to_docx(SAMPLE_MD, title="审查意见书")
    assert isinstance(data, bytes)
    assert data[:2] == b"PK"  # docx 即 zip
    doc = Document(BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "一、文件基本情况" in text
    assert "二、审查依据" in text
    assert "《中小企业促进法》第十条" in text


def test_markdown_to_docx_preserves_bold() -> None:
    """**加粗** 行内标记还原为 bold run。"""
    from docx import Document

    data = markdown_to_docx("本文件为 **重要文件** 的说明。")
    doc = Document(BytesIO(data))
    bold_texts = [
        run.text
        for p in doc.paragraphs
        for run in p.runs
        if run.bold
    ]
    assert "重要文件" in bold_texts


def test_markdown_to_docx_renders_table() -> None:
    """管道表格转为 Word 表格，含表头行与数据行。"""
    from docx import Document

    data = markdown_to_docx(SAMPLE_MD)
    doc = Document(BytesIO(data))
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 3  # 表头 + 2 数据行
    assert table.rows[0].cells[0].text == "项目"
    assert table.rows[1].cells[1].text == "PASS"


def test_markdown_to_docx_empty_markdown() -> None:
    """空 Markdown 不报错，至少含标题段落。"""
    from docx import Document

    data = markdown_to_docx("", title="审查意见书")
    doc = Document(BytesIO(data))
    assert any(p.text == "审查意见书" for p in doc.paragraphs)


# ---------- HTTP 层 ----------


class _FakeResult:
    def __init__(self, value: object) -> None:
        self._value = value

    def scalar_one_or_none(self) -> object:
        return self._value


class _QueueSession:
    """按调用顺序依次返回预设结果的假 Session（不触 DB）。"""

    def __init__(self, results: list[object]) -> None:
        self._results = list(results)
        self._i = 0

    async def execute(self, _stmt: object) -> _FakeResult:
        value = self._results[self._i]
        self._i += 1
        return _FakeResult(value)


def _make_task_and_report() -> tuple[SimpleNamespace, SimpleNamespace]:
    task_id = uuid4()
    org_id = uuid4()
    user_id = uuid4()
    task = SimpleNamespace(
        id=task_id,
        status="done",
        current_node="report_generation",
        organization_id=org_id,
        submitter_id=user_id,
    )
    rr = SimpleNamespace(
        output_json={"report_markdown": SAMPLE_MD},
        risks=[],
        evidences=[],
        agent_name="report_generation",
    )
    return task, rr


def _client(results: list[object]) -> AsyncClient:
    app = FastAPI()
    app.include_router(tasks_router)
    session = _QueueSession(results)

    async def _override_db():
        yield session

    async def _override_ctx() -> TenantContext:
        return TenantContext(
            user=CorePrincipal(
                id=uuid4(),
                tenant_id=None,
                email="admin@example.com",
                full_name="Admin",
                role="admin",
                is_active=True,
            ),
            org=None,
            org_id=None,
            is_team=False,
            is_super_admin=False,
        )

    app.dependency_overrides[get_db] = _override_db
    app.dependency_overrides[get_tenant_context] = _override_ctx

    @app.exception_handler(AppError)
    async def _app_error_handler(_req: Request, exc: AppError) -> JSONResponse:
        return JSONResponse(status_code=exc.http_status, content=error_response(exc))

    return AsyncClient(transport=ASGITransport(app=app), base_url="http://test")


async def test_export_endpoint_returns_docx() -> None:
    """用例①：报告存在 → 200，Content-Type 为 docx，body 为合法 docx。"""
    task, rr = _make_task_and_report()
    async with _client([task, rr]) as client:
        resp = await client.get(f"/tasks/{task.id}/report/export")
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert "attachment" in resp.headers["content-disposition"]
    assert "filename*=UTF-8''" in resp.headers["content-disposition"]
    assert resp.content[:2] == b"PK"


async def test_export_endpoint_custom_filename() -> None:
    """用例②：filename 参数生效 → Content-Disposition 携带编码后的中文名。"""
    task, rr = _make_task_and_report()
    async with _client([task, rr]) as client:
        resp = await client.get(
            f"/tasks/{task.id}/report/export", params={"filename": "XX县意见审查意见书"}
        )
    assert resp.status_code == 200
    from urllib.parse import quote

    assert quote("XX县意见审查意见书.docx") in resp.headers["content-disposition"]


async def test_export_endpoint_report_not_ready_404() -> None:
    """用例③：报告未生成（ReviewResult 缺失）→ 404。"""
    task, _ = _make_task_and_report()
    async with _client([task, None]) as client:
        resp = await client.get(f"/tasks/{task.id}/report/export")
    assert resp.status_code == 404


async def test_export_endpoint_invalid_format_422() -> None:
    """用例④：不支持的 format → 422（仅 docx）。"""
    task, rr = _make_task_and_report()
    async with _client([task, rr]) as client:
        resp = await client.get(
            f"/tasks/{task.id}/report/export", params={"format": "pdf"}
        )
    assert resp.status_code == 422
