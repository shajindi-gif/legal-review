"""规范性文件合法性审核增强测试（PHASE 1-5）。

四层：
1. 规则包层：注册表加载 / 按地区解析 / 规则字段完整性。
2. 报告层：10 段输出标准 + 4 类结论（确定性，不依赖 LLM）。
3. 导出层：意见书 DOCX / 报告 PDF / 条款问题清单 XLSX。
4. HTTP 层：/gov/rule-packs(+resolve) 与 /tasks/{id}/gov-report、gov-export。
"""

from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

import pytest
from fastapi import FastAPI, Request
from fastapi.responses import JSONResponse
from httpx import ASGITransport, AsyncClient

from app.api.deps import TenantContext, get_current_user, get_db, get_tenant_context
from app.api.v1.gov import router as gov_router
from app.api.v1.tasks import router as tasks_router
from app.core.constants import GovReviewConclusion
from app.core.errors import AppError, NotFoundError, error_response
from app.schemas.gov_review import GOV_REPORT_SECTION_TITLES, GovReviewReport
from app.services.core_auth_client import CorePrincipal
from app.services.gov_export import (
    render_findings_xlsx,
    render_opinion_docx,
    render_report_pdf,
    report_to_markdown,
)
from app.services.gov_review_report import build_gov_report
from app.services.gov_rules import (
    GovRule,
    GovRulePack,
    get_gov_rule_pack_manager,
)


# ============== 测试替身 ==============
def _rule(rule_id: str, dimension: str, severity: str, material: str) -> GovRule:
    return GovRule(
        rule_id=rule_id,
        dimension=dimension,
        title=f"{dimension} 维度测试规则",
        risk_type=f"{dimension} 风险",
        severity=severity,
        law_basis=f"《测试{rule_id}法》",
        region="全国",
        condition="核查条件",
        evidence_required=[material],
    )


def _packs() -> list[GovRulePack]:
    return [
        GovRulePack(
            pack_id="test_national",
            name="测试国家通用规则包",
            scope="national",
            region="全国",
            description="",
            legal_review_status="pending_legal_review",
            rules=[
                _rule("T-A-01", "A", "critical", "职权依据文件"),
                _rule("T-E-01", "E", "medium", "征求意见记录"),
            ],
        )
    ]


def _result(
    *,
    risks: list[dict] | None = None,
    evidences: list[dict] | None = None,
    output_json: dict | None = None,
) -> SimpleNamespace:
    return SimpleNamespace(
        risks=risks or [],
        evidences=evidences or [],
        output_json=output_json or {},
        node_status="pass",
        confidence=0.8,
    )


def _critical_risk() -> dict:
    return {
        "dimension": "authority",
        "risk_type": "违法设定行政许可",
        "severity": "critical",
        "paragraph_id": "p3",
        "evidence": {
            "law_name": "行政许可法",
            "article": "第十五条",
            "original_text": "……",
            "explanation": "与上位法冲突",
        },
        "confidence": 0.9,
        "suggestion": "删除该许可条款",
    }


# ============== 1. 规则包层 ==============
def test_rule_packs_loaded_and_national_first() -> None:
    packs = get_gov_rule_pack_manager().list_packs()
    ids = [p.pack_id for p in packs]
    assert ids[0] == "national_general"
    assert {"national_general", "shanghai", "xuhui"} <= set(ids)


def test_rule_pack_get_unknown_raises() -> None:
    with pytest.raises(NotFoundError):
        get_gov_rule_pack_manager().get_pack("not-exist")


def test_resolve_for_region_district() -> None:
    packs = get_gov_rule_pack_manager().resolve_for_region("上海市徐汇区")
    ids = {p.pack_id for p in packs}
    assert ids == {"national_general", "shanghai", "xuhui"}


def test_resolve_for_region_city_only() -> None:
    packs = get_gov_rule_pack_manager().resolve_for_region("上海市")
    ids = {p.pack_id for p in packs}
    assert ids == {"national_general", "shanghai"}


def test_resolve_for_region_none_returns_national_only() -> None:
    packs = get_gov_rule_pack_manager().resolve_for_region(None)
    assert [p.pack_id for p in packs] == ["national_general"]


def test_all_rules_have_required_fields() -> None:
    """每条规则必须有法律依据 / 适用地区 / 审核条件 / 风险类型，维度须属 A~J。"""
    rules = get_gov_rule_pack_manager().rules_for_region("上海市徐汇区")
    assert rules, "规则包应展开出规则"
    valid_dims = set("ABCDEFGHIJ")
    for rule in rules:
        assert rule.dimension in valid_dims
        assert rule.law_basis.strip()
        assert rule.region.strip()
        assert rule.condition.strip()
        assert rule.risk_type.strip()


# ============== 2. 报告层 ==============
def test_report_has_ten_sections() -> None:
    report = build_gov_report(
        results=[_result()], packs=_packs(), document=None, task=None
    )
    assert isinstance(report, GovReviewReport)
    assert len(GOV_REPORT_SECTION_TITLES) == 10
    for key, _title in GOV_REPORT_SECTION_TITLES:
        assert hasattr(report, key)


def test_conclusion_major_risk_for_critical_finding() -> None:
    report = build_gov_report(
        results=[
            _result(
                risks=[_critical_risk()],
                evidences=[{"law_name": "行政许可法", "article": "第十五条"}],
            )
        ],
        packs=_packs(),
        task=None,
    )
    assert report.conclusion == GovReviewConclusion.MAJOR_RISK
    assert report.findings_count == 1
    assert report.findings[0].dimension == "C"  # 关键词“许可”→ C
    assert report.overall_risk == "critical"
    assert report.suggestions[0].suggestion == "删除该许可条款"
    assert report.clause_locations[0].clause_location == "p3"


def test_conclusion_needs_revision_for_medium_finding() -> None:
    medium = dict(_critical_risk(), severity="medium", risk_type="程序瑕疵")
    report = build_gov_report(
        results=[_result(risks=[medium])], packs=_packs(), task=None
    )
    assert report.conclusion == GovReviewConclusion.NEEDS_REVISION
    assert report.findings[0].dimension == "E"  # 关键词“程序”→ E


def test_conclusion_insufficient_when_no_findings_and_materials_missing() -> None:
    report = build_gov_report(
        results=[_result()],
        packs=_packs(),
        document=SimpleNamespace(original_name="x.docx", parsed_json={}),
        task=None,
    )
    assert report.conclusion == GovReviewConclusion.INSUFFICIENT_BASIS
    missing = {m.material for m in report.missing_materials}
    assert missing == {"职权依据文件", "征求意见记录"}


def test_conclusion_no_obvious_violation_when_materials_complete() -> None:
    doc = SimpleNamespace(
        original_name="x.docx",
        parsed_json={"submitted_materials": ["职权依据文件", "征求意见记录"]},
    )
    report = build_gov_report(results=[_result()], packs=_packs(), document=doc)
    assert report.conclusion == GovReviewConclusion.NO_OBVIOUS_VIOLATION
    assert report.missing_materials == []


def test_pending_items_include_pack_legal_review() -> None:
    report = build_gov_report(results=[_result()], packs=_packs())
    reasons = [p.reason for p in report.pending_human_review]
    assert any("法律专业人员审核" in r for r in reasons)


def test_procedure_missing_marked_insufficient_not_illegal() -> None:
    """硬约束：程序材料缺失只标记「材料不足」，不认定程序违法。"""
    report = build_gov_report(results=[_result()], packs=_packs())
    for m in report.missing_materials:
        assert "材料不足" in m.reason
        # 否定式表述「不据此认定程序违法」== 未作程序违法认定（不编造结论）
        assert "不据此认定程序违法" in m.reason
    assert report.conclusion == GovReviewConclusion.INSUFFICIENT_BASIS


# ============== 3. 导出层 ==============
def _sample_report() -> GovReviewReport:
    return build_gov_report(
        results=[
            _result(
                risks=[_critical_risk()],
                evidences=[{"law_name": "行政许可法", "article": "第十五条"}],
            )
        ],
        packs=_packs(),
        document=SimpleNamespace(original_name="草案.docx", parsed_json={}),
        task=SimpleNamespace(id=uuid4()),
    )


def test_report_to_markdown_has_all_titles() -> None:
    md = report_to_markdown(_sample_report())
    for _key, title in GOV_REPORT_SECTION_TITLES:
        assert title in md


def test_render_opinion_docx_is_valid() -> None:
    from docx import Document

    data = render_opinion_docx(_sample_report())
    assert data[:2] == b"PK"
    doc = Document(BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "规范性文件合法性审核意见书" in text


def test_render_report_pdf_is_valid() -> None:
    data = render_report_pdf(_sample_report())
    assert data[:5] == b"%PDF-"
    assert len(data) > 1000


def test_render_findings_xlsx_is_valid() -> None:
    from openpyxl import load_workbook

    data = render_findings_xlsx(_sample_report())
    assert data[:2] == b"PK"
    wb = load_workbook(BytesIO(data))
    assert wb.sheetnames == ["条款问题清单", "材料缺失清单"]
    ws = wb["条款问题清单"]
    assert ws.cell(row=1, column=3).value == "问题"
    assert ws.max_row == 2  # 表头 + 1 条问题


# ============== 4. HTTP 层 ==============
class _ScalarResult:
    def __init__(self, value: object) -> None:
        self._value = value

    def scalar_one_or_none(self) -> object:
        return self._value


class _ScalarsResult:
    def __init__(self, values: list[object]) -> None:
        self._values = values

    def scalars(self) -> _ScalarsResult:
        return self

    def all(self) -> list[object]:
        return list(self._values)


class _QueueSession:
    def __init__(self, results: list[object]) -> None:
        self._results = list(results)
        self._i = 0

    async def execute(self, _stmt: object) -> object:
        value = self._results[self._i]
        self._i += 1
        return value


def _admin_ctx() -> TenantContext:
    return TenantContext(
        user=CorePrincipal(
            id=uuid4(),
            tenant_id=None,
            email="admin@example.com",
            full_name="Admin",
            role="admin",
            is_active=True,
        ),
        org=None,
        org_id=None,
        is_team=False,
        is_super_admin=False,
    )


def _client(
    *,
    include_tasks: bool,
    db: object | None = None,
) -> AsyncClient:
    app = FastAPI()
    app.include_router(gov_router)
    if include_tasks:
        app.include_router(tasks_router)

    if db is not None:

        async def _override_db() -> object:
            yield db

        app.dependency_overrides[get_db] = _override_db

    async def _override_user() -> CorePrincipal:
        return _admin_ctx().user

    async def _override_ctx() -> TenantContext:
        return _admin_ctx()

    app.dependency_overrides[get_current_user] = _override_user
    app.dependency_overrides[get_tenant_context] = _override_ctx

    @app.exception_handler(AppError)
    async def _app_error_handler(_req: Request, exc: AppError) -> JSONResponse:
        return JSONResponse(status_code=exc.http_status, content=error_response(exc))

    return AsyncClient(transport=ASGITransport(app=app), base_url="http://test")


async def test_api_list_rule_packs() -> None:
    async with _client(include_tasks=False) as client:
        resp = await client.get("/gov/rule-packs")
    assert resp.status_code == 200
    body = resp.json()
    assert len(body) >= 3
    assert body[0]["pack_id"] == "national_general"
    assert "rules" in body[0]


async def test_api_resolve_rule_packs_district() -> None:
    async with _client(include_tasks=False) as client:
        resp = await client.get(
            "/gov/rule-packs/resolve", params={"region": "上海市徐汇区"}
        )
    assert resp.status_code == 200
    ids = {p["pack_id"] for p in resp.json()}
    assert ids == {"national_general", "shanghai", "xuhui"}


async def test_api_get_rule_pack() -> None:
    async with _client(include_tasks=False) as client:
        resp = await client.get("/gov/rule-packs/shanghai")
    assert resp.status_code == 200
    assert resp.json()["pack_id"] == "shanghai"


async def test_api_gov_report_returns_ten_sections() -> None:
    task_id = uuid4()
    task = SimpleNamespace(
        id=task_id,
        status="done",
        current_node="report_generation",
        organization_id=uuid4(),
        submitter_id=uuid4(),
        title="草案",
        metadata_={},
    )
    doc = SimpleNamespace(id=uuid4(), parsed_json={}, original_name="草案.docx")
    session = _QueueSession(
        [
            _ScalarResult(task),
            _ScalarResult(doc),
            _ScalarsResult(
                [
                    _result(
                        risks=[_critical_risk()],
                        evidences=[{"law_name": "行政许可法", "article": "第十五条"}],
                    )
                ]
            ),
        ]
    )
    async with _client(include_tasks=True, db=session) as client:
        resp = await client.get(f"/tasks/{task_id}/gov-report")
    assert resp.status_code == 200, resp.text
    body = resp.json()
    for key, _title in GOV_REPORT_SECTION_TITLES:
        assert key in body
    assert body["conclusion"] == "major_risk"
    assert body["findings_count"] == 1


async def test_api_gov_export_xlsx() -> None:
    task_id = uuid4()
    task = SimpleNamespace(
        id=task_id,
        status="done",
        current_node="report_generation",
        organization_id=uuid4(),
        submitter_id=uuid4(),
        title="草案",
        metadata_={},
    )
    session = _QueueSession(
        [
            _ScalarResult(task),
            _ScalarResult(SimpleNamespace(parsed_json={}, original_name="草案.docx")),
            _ScalarsResult([_result(risks=[_critical_risk()])]),
        ]
    )
    async with _client(include_tasks=True, db=session) as client:
        resp = await client.get(
            f"/tasks/{task_id}/gov-export", params={"format": "xlsx"}
        )
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert resp.content[:2] == b"PK"


async def test_api_gov_export_invalid_format_422() -> None:
    task_id = uuid4()
    session = _QueueSession([])
    async with _client(include_tasks=True, db=session) as client:
        resp = await client.get(
            f"/tasks/{task_id}/gov-export", params={"format": "doc"}
        )
    assert resp.status_code == 422
