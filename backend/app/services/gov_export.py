"""规范性文件合法性审核交付物导出（PHASE 1-5）。

三种交付物（指令第六条）：
- 规范性文件合法性审核意见书 DOCX   → render_opinion_docx（复用 report_export）
- 规范性文件审核报告 PDF            → render_report_pdf（HTML → PDF）
- 条款问题清单 XLSX                 → render_findings_xlsx

均为纯本地转换（不依赖 LLM / 网络），便于离线与测试复现。
"""

from __future__ import annotations

import html
from io import BytesIO

from app.schemas.gov_review import (
    GOV_REPORT_SECTION_TITLES,
    GovReviewReport,
)
from app.services.report_export import markdown_to_docx

DOCX_MIME = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
PDF_MIME = "application/pdf"
XLSX_MIME = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

EXPORT_FORMATS: tuple[str, ...] = ("docx", "pdf", "xlsx")


# ============== Markdown（DOCX 输入）==============
def report_to_markdown(report: GovReviewReport) -> str:
    """把 10 段报告渲染为 Markdown（供意见书 DOCX 复用）。"""
    lines: list[str] = ["# 规范性文件合法性审核报告（AI 辅助）", ""]
    lines.append(f"> {report.disclaimer}")
    lines.append("")

    titles = dict(GOV_REPORT_SECTION_TITLES)

    basic = report.basic_info
    lines.append(f"## 一、{titles['basic_info']}")
    lines.append(f"- 文件名称：{basic.file_name or '未识别'}")
    lines.append(f"- 制定机关：{basic.issuing_authority or '未识别'}")
    lines.append(f"- 发文形式：{basic.doc_form or '未识别'}")
    lines.append(f"- 文号：{basic.doc_number or '未识别'}")
    lines.append(f"- 适用区域：{basic.region or '未识别'}")
    if basic.involved_matters:
        lines.append(f"- 涉及事项：{'、'.join(basic.involved_matters)}")
    lines.append("")

    scope = report.review_scope
    lines.append(f"## 二、{titles['review_scope']}")
    lines.append(f"- 适用区域：{scope.region or '未识别'}")
    lines.append(f"- 适用规则包：{'、'.join(scope.pack_ids) or '无'}")
    lines.append(f"- 审核维度：{'、'.join(scope.dimension_labels) or '无'}")
    lines.append("")

    basis = report.review_basis
    lines.append(f"## 三、{titles['review_basis']}")
    lines.append("### 规则包法律依据")
    if basis.rule_basis:
        lines.extend(f"- {b}" for b in basis.rule_basis)
    else:
        lines.append("- 无")
    lines.append("### 文件引用依据")
    if basis.cited_basis:
        lines.extend(f"- {b}" for b in basis.cited_basis)
    else:
        lines.append("- 无")
    lines.append("")

    lines.append(f"## 四、{titles['findings']}")
    if report.findings:
        lines.append("| 编号 | 维度 | 问题 | 条款位置 | 风险等级 | 建议 |")
        lines.append("| --- | --- | --- | --- | --- | --- |")
        for f in report.findings:
            lines.append(
                f"| {f.id} | {f.dimension} | {f.title} | "
                f"{f.clause_location or '-'} | {f.risk_level} | "
                f"{f.suggestion or '-'} |"
            )
    else:
        lines.append("- 未发现明显违法问题（以规则包为准，需人工复核）")
    lines.append("")

    lines.append(f"## 五、{titles['clause_locations']}")
    if report.clause_locations:
        for c in report.clause_locations:
            lines.append(f"- {c.finding_id}（维度 {c.dimension}）：{c.clause_location or '未定位'}")
    else:
        lines.append("- 无")
    lines.append("")

    lines.append(f"## 六、{titles['risk_levels']}")
    lines.append("| 等级 | 数量 |")
    lines.append("| --- | --- |")
    for r in report.risk_levels:
        lines.append(f"| {r.label} | {r.count} |")
    lines.append(f"\n**审核结论：{report.conclusion_label}**")
    lines.append("")

    lines.append(f"## 七、{titles['suggestions']}")
    if report.suggestions:
        for s in report.suggestions:
            lines.append(f"- {s.finding_id}（维度 {s.dimension}）：{s.suggestion}")
    else:
        lines.append("- 无")
    lines.append("")

    lines.append(f"## 八、{titles['missing_materials']}")
    if report.missing_materials:
        lines.extend(
            f"- {m.material}（规则 {m.rule_id or '-'}，维度 {m.dimension or '-'}）：{m.reason}"
            for m in report.missing_materials
        )
    else:
        lines.append("- 送审材料已齐备")
    lines.append("")

    lines.append(f"## 九、{titles['pending_human_review']}")
    if report.pending_human_review:
        lines.extend(f"- {p.item}（{p.reason}）" for p in report.pending_human_review)
    else:
        lines.append("- 无")
    lines.append("")

    lines.append(f"## 十、{titles['opinion_draft']}")
    lines.append(report.opinion_draft)
    return "\n".join(lines)


# ============== DOCX ==============
def render_opinion_docx(report: GovReviewReport) -> bytes:
    """规范性文件合法性审核意见书 DOCX。"""
    return markdown_to_docx(
        report_to_markdown(report), title="规范性文件合法性审核意见书"
    )


# ============== PDF ==============
def _html(report: GovReviewReport) -> str:
    esc = html.escape

    rows_findings = "".join(
        "<tr>"
        f"<td>{esc(f.id)}</td><td>{esc(f.dimension)}</td>"
        f"<td>{esc(f.title)}</td><td>{esc(f.clause_location or '-')}</td>"
        f"<td>{esc(str(f.risk_level))}</td><td>{esc(f.suggestion or '-')}</td>"
        "</tr>"
        for f in report.findings
    ) or "<tr><td colspan='6'>未发现明显违法问题（需人工复核）</td></tr>"

    rows_risk = "".join(
        f"<tr><td>{esc(r.label)}</td><td>{r.count}</td></tr>"
        for r in report.risk_levels
    )

    rows_missing = "".join(
        f"<tr><td>{esc(m.material)}</td><td>{esc(m.dimension or '-')}</td>"
        f"<td>{esc(m.rule_id or '-')}</td><td>{esc(m.reason)}</td></tr>"
        for m in report.missing_materials
    ) or "<tr><td colspan='4'>送审材料已齐备</td></tr>"

    rows_pending = "".join(
        f"<li>{esc(p.item)}（{esc(p.reason)}）</li>"
        for p in report.pending_human_review
    ) or "<li>无</li>"

    basic = report.basic_info
    return f"""<!DOCTYPE html>
<html lang="zh-CN"><head><meta charset="utf-8">
<style>
  @page {{ size: A4; margin: 2cm; }}
  body {{ font-family: "Songti SC","STSong","Noto Sans CJK SC","Source Han Sans SC",sans-serif;
          font-size: 11pt; color: #1a1a1a; line-height: 1.7; }}
  h1 {{ font-size: 18pt; text-align: center; }}
  h2 {{ font-size: 13pt; border-bottom: 1px solid #ccc; padding-bottom: 4px; }}
  table {{ border-collapse: collapse; width: 100%; margin: 8px 0; }}
  th, td {{ border: 1px solid #bbb; padding: 4px 6px; font-size: 10pt; text-align: left; }}
  th {{ background: #f2f2f2; }}
  .disclaimer {{ color: #666; font-size: 9.5pt; }}
</style></head><body>
<h1>规范性文件合法性审核报告</h1>
<p class="disclaimer">{esc(report.disclaimer)}</p>

<h2>一、基本信息</h2>
<table>
  <tr><th>文件名称</th><td>{esc(basic.file_name or '未识别')}</td></tr>
  <tr><th>制定机关</th><td>{esc(basic.issuing_authority or '未识别')}</td></tr>
  <tr><th>发文形式</th><td>{esc(basic.doc_form or '未识别')}</td></tr>
  <tr><th>文号</th><td>{esc(basic.doc_number or '未识别')}</td></tr>
  <tr><th>适用区域</th><td>{esc(basic.region or '未识别')}</td></tr>
</table>

<h2>二、审核范围</h2>
<p>适用区域：{esc(report.review_scope.region or '未识别')}；
适用规则包：{esc('、'.join(report.review_scope.pack_ids) or '无')}；
审核维度：{esc('、'.join(report.review_scope.dimension_labels) or '无')}</p>

<h2>三、审核依据</h2>
<p>规则包法律依据：{esc('、'.join(report.review_basis.rule_basis) or '无')}</p>
<p>文件引用依据：{esc('、'.join(report.review_basis.cited_basis) or '无')}</p>

<h2>四、发现问题</h2>
<table>
  <tr><th>编号</th><th>维度</th><th>问题</th><th>条款位置</th><th>风险等级</th><th>修改建议</th></tr>
  {rows_findings}
</table>

<h2>六、风险等级与结论</h2>
<table><tr><th>等级</th><th>数量</th></tr>{rows_risk}</table>
<p><strong>审核结论：{esc(report.conclusion_label)}</strong></p>

<h2>八、材料缺失清单</h2>
<table><tr><th>材料</th><th>维度</th><th>关联规则</th><th>说明</th></tr>{rows_missing}</table>

<h2>九、待人工核实事项</h2>
<ul>{rows_pending}</ul>
</body></html>"""


def render_report_pdf(report: GovReviewReport) -> bytes:
    """规范性文件审核报告 PDF（HTML → PDF）。"""
    from typing import cast

    from weasyprint import HTML

    pdf = cast(bytes, HTML(string=_html(report)).write_pdf())
    return pdf


# ============== XLSX ==============
def render_findings_xlsx(report: GovReviewReport) -> bytes:
    """条款问题清单 XLSX（含「条款问题清单」「材料缺失清单」两个工作表）。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4A5568")
    wrap = Alignment(wrap_text=True, vertical="top")

    ws = wb.active
    ws.title = "条款问题清单"
    headers = [
        "序号",
        "审核维度",
        "问题",
        "条款位置",
        "风险等级",
        "法律依据",
        "风险原因",
        "修改建议",
        "需人工复核",
    ]
    ws.append(headers)
    for idx, f in enumerate(report.findings, 1):
        basis = f.law_basis or ""
        if f.law_article:
            basis = f"{basis}{f.law_article}"
        ws.append([
            idx,
            f.dimension,
            f.title,
            f.clause_location or "",
            str(f.risk_level),
            basis,
            f.explanation or "",
            f.suggestion or "",
            "是" if f.needs_human_review else "否",
        ])

    ws2 = wb.create_sheet("材料缺失清单")
    ws2.append(["序号", "材料名称", "关联维度", "关联规则", "说明"])
    for idx, m in enumerate(report.missing_materials, 1):
        ws2.append([idx, m.material, m.dimension or "", m.rule_id or "", m.reason])

    for sheet, widths in (
        (ws, [6, 10, 34, 14, 10, 22, 34, 34, 12]),
        (ws2, [6, 30, 12, 14, 40]),
    ):
        for i, width in enumerate(widths, 1):
            sheet.column_dimensions[sheet.cell(row=1, column=i).column_letter].width = width
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = wrap
        sheet.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


__all__ = [
    "DOCX_MIME",
    "EXPORT_FORMATS",
    "PDF_MIME",
    "XLSX_MIME",
    "render_findings_xlsx",
    "render_opinion_docx",
    "render_report_pdf",
    "report_to_markdown",
]
