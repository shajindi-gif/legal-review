"""报告导出：Markdown → Word(.docx)。

用于「审查意见书」的可归档交付物导出（R1-1）。
不依赖 LLM / 网络，纯本地转换，便于离线与测试复现。

支持的 Markdown 子集（覆盖 report_generation 提示词输出）：
- ATX 标题 `#` ~ `######`
- 无序列表 `-` / `*` / `+`
- 有序列表 `1.` / `1、`
- 引用块 `>`
- 围栏代码块 ```（原样成段，等宽字体）
- 水平分割线 `---` / `***`
- 简单管道表格 `| a | b |`
- 行内 **加粗** / *斜体*
"""

from __future__ import annotations

import re
from io import BytesIO

_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_ULIST_RE = re.compile(r"^[-*+]\s+(.*)$")
_OLIST_RE = re.compile(r"^(\d+)[.、]\s+(.*)$")
_HR_RE = re.compile(r"^(-{3,}|\*{3,}|_{3,})$")
_TABLE_SEP_RE = re.compile(r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)+\|?$")
_FENCE_RE = re.compile(r"^(`{3,}|~{3,})")


def _add_runs(paragraph: object, text: str) -> None:
    """把含 **加粗** / *斜体* 的行内文本写入段落。"""
    # 先按加粗切分，再在每段内处理斜体
    for bold_chunk in re.split(r"(\*\*.+?\*\*)", text):
        if not bold_chunk:
            continue
        if bold_chunk.startswith("**") and bold_chunk.endswith("**") and len(bold_chunk) > 4:
            run = paragraph.add_run(bold_chunk[2:-2])  # type: ignore[attr-defined]
            run.bold = True
            continue
        for italic_chunk in re.split(r"(\*[^*]+?\*)", bold_chunk):
            if not italic_chunk:
                continue
            if (
                italic_chunk.startswith("*")
                and italic_chunk.endswith("*")
                and len(italic_chunk) > 2
            ):
                run = paragraph.add_run(italic_chunk[1:-1])  # type: ignore[attr-defined]
                run.italic = True
            else:
                paragraph.add_run(italic_chunk)  # type: ignore[attr-defined]


def _add_table(doc: object, rows: list[list[str]]) -> None:
    """写入管道表格（首行表头）。"""
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=ncols)  # type: ignore[attr-defined]
    table.style = "Light Grid Accent 1"
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j in range(ncols):
            text = row[j] if j < len(row) else ""
            cells[j].text = text
            if i == 0:
                for p in cells[j].paragraphs:
                    for run in p.runs:
                        run.bold = True


def _split_table_row(line: str) -> list[str]:
    inner = line.strip()
    if inner.startswith("|"):
        inner = inner[1:]
    if inner.endswith("|"):
        inner = inner[:-1]
    return [c.strip() for c in inner.split("|")]


def markdown_to_docx(markdown: str, *, title: str = "审查意见书") -> bytes:
    """把报告 Markdown 转换为 Word 文档字节。"""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(10.5)

    doc.add_heading(title, level=0)

    lines = (markdown or "").splitlines()
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i].rstrip()
        stripped = line.strip()

        # 围栏代码块
        if _FENCE_RE.match(stripped):
            i += 1
            code_lines: list[str] = []
            while i < n and not _FENCE_RE.match(lines[i].strip()):
                code_lines.append(lines[i])
                i += 1
            i += 1  # 跳过闭合围栏
            p = doc.add_paragraph("\n".join(code_lines))
            p.style = doc.styles["No Spacing"]
            for run in p.runs:
                run.font.name = "Consolas"
            continue

        if not stripped:
            i += 1
            continue

        # 标题
        m = _HEADING_RE.match(stripped)
        if m:
            level = min(len(m.group(1)), 4)
            doc.add_heading(m.group(2).strip(), level=level)
            i += 1
            continue

        # 水平分割线
        if _HR_RE.match(stripped):
            doc.add_paragraph("─" * 30)
            i += 1
            continue

        # 表格：当前行含 | 且下一行是分隔行
        if "|" in stripped and i + 1 < n and _TABLE_SEP_RE.match(lines[i + 1].strip()):
            table_rows = [_split_table_row(stripped)]
            i += 2  # 跳过分隔行
            while i < n and "|" in lines[i] and lines[i].strip():
                table_rows.append(_split_table_row(lines[i]))
                i += 1
            _add_table(doc, table_rows)
            continue

        # 引用块
        if stripped.startswith(">"):
            text = stripped.lstrip(">").strip()
            p = doc.add_paragraph()
            _add_runs(p, text)
            p.paragraph_format.left_indent = Pt(24)
            i += 1
            continue

        # 有序列表
        m = _OLIST_RE.match(stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            _add_runs(p, m.group(2))
            i += 1
            continue

        # 无序列表
        m = _ULIST_RE.match(stripped)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            _add_runs(p, m.group(1))
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        _add_runs(p, stripped)
        i += 1

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
